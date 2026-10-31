import pandas as pd
import openpyxl
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os

# Set paths
project1_dir = os.path.dirname(os.path.abspath(__file__))
cleaned_data_path = os.path.join(project1_dir, "Day12_datacleaning", "cleaned_online_retail.csv")
output_excel_path = os.path.join(project1_dir, "project1_dashboard.xlsx")

print("Loading cleaned dataset...")
df = pd.read_csv(cleaned_data_path)

# Ensure datetime and total price columns exist
df['InvoiceDate'] = pd.to_datetime(df['InvoiceDate'])
if 'TotalPrice' not in df.columns:
    df['TotalPrice'] = df['Quantity'] * df['Price']

# 1. Calculate Core KPIs
total_revenue = df['TotalPrice'].sum()
total_orders = df['Invoice'].nunique()
total_customers = df['Customer ID'].dropna().nunique()
aov = total_revenue / total_orders if total_orders > 0 else 0

# 2. Aggregations for Charts
# Monthly Revenue
df['YearMonth'] = df['InvoiceDate'].dt.to_period('M').astype(str)
monthly_df = df.groupby('YearMonth')['TotalPrice'].sum().reset_index()
monthly_df.columns = ['Month', 'Revenue']

# Top 5 Products by Revenue
top_products_df = df.groupby('Description')['TotalPrice'].sum().reset_index()
top_products_df = top_products_df.sort_values(by='TotalPrice', ascending=False).head(5)
top_products_df.columns = ['Product', 'Revenue']

# Top 5 Countries by Revenue
top_countries_df = df.groupby('Country')['TotalPrice'].sum().reset_index()
top_countries_df = top_countries_df.sort_values(by='TotalPrice', ascending=False).head(5)
top_countries_df.columns = ['Country', 'Revenue']

# 3. Create Excel Workbook using OpenPyXL
wb = openpyxl.Workbook()

# Setup Sheets
ws_dash = wb.active
ws_dash.title = "Dashboard"
ws_data = wb.create_sheet(title="Data_Summary")

# Styling Definitions
font_title = Font(name="Calibri", size=18, bold=True, color="1F497D")
font_subtitle = Font(name="Calibri", size=11, italic=True, color="595959")
font_kpi_label = Font(name="Calibri", size=10, bold=True, color="595959")
font_kpi_val = Font(name="Calibri", size=18, bold=True, color="1F497D")
font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

fill_header = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
fill_kpi = PatternFill(start_color="F2F4F7", end_color="F2F4F7", fill_type="solid")

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

# --- POPULATE DASHBOARD SHEET ---
# Title Header
ws_dash['B2'] = "Online Retail II Performance Dashboard"
ws_dash['B2'].font = font_title
ws_dash['B3'] = "Project 1 Executive Summary | Dimple Singh K (U19XF24S0021)"
ws_dash['B3'].font = font_subtitle

# KPI Card Helper Function
def create_kpi_card(ws, start_col, label, value, number_format=None):
    col1, col2 = start_col, chr(ord(start_col) + 1)
    
    # Merge cells for Label and Value
    ws.merge_cells(f"{col1}5:{col2}5")
    ws.merge_cells(f"{col1}6:{col2}7")
    
    lbl_cell = ws[f"{col1}5"]
    val_cell = ws[f"{col1}6"]
    
    lbl_cell.value = label
    lbl_cell.font = font_kpi_label
    lbl_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    val_cell.value = value
    val_cell.font = font_kpi_val
    val_cell.alignment = Alignment(horizontal="center", vertical="center")
    if number_format:
        val_cell.number_format = number_format

    # Apply background fill and borders
    for row in range(5, 8):
        for c in [col1, col2]:
            cell = ws[f"{c}{row}"]
            cell.fill = fill_kpi
            cell.border = thin_border

create_kpi_card(ws_dash, 'B', "TOTAL REVENUE", total_revenue, "₹#,##0.00")
create_kpi_card(ws_dash, 'E', "TOTAL ORDERS", total_orders, "#,##0")
create_kpi_card(ws_dash, 'H', "AVG ORDER VALUE", aov, "₹#,##0.00")
create_kpi_card(ws_dash, 'K', "TOTAL CUSTOMERS", total_customers, "#,##0")

# --- POPULATE DATA_SUMMARY SHEET ---
ws_data.append(["Month", "Monthly Revenue"])
for r in monthly_df.itertuples(index=False):
    ws_data.append(list(r))

ws_data.append([])
ws_data.append(["Product", "Top Product Revenue"])
for r in top_products_df.itertuples(index=False):
    ws_data.append(list(r))

ws_data.append([])
ws_data.append(["Country", "Country Revenue"])
for r in top_countries_df.itertuples(index=False):
    ws_data.append(list(r))

# --- ADD CHARTS TO DASHBOARD SHEET ---
# 1. Monthly Revenue Line Chart
chart_line = LineChart()
chart_line.title = "Monthly Revenue Trend"
chart_line.style = 13
chart_line.y_axis.title = "Revenue"
chart_line.x_axis.title = "Month"

data_line = Reference(ws_data, min_col=2, min_row=1, max_row=len(monthly_df)+1)
cats_line = Reference(ws_data, min_col=1, min_row=2, max_row=len(monthly_df)+1)
chart_line.add_data(data_line, titles_from_data=True)
chart_line.set_categories(cats_line)
chart_line.width = 16
chart_line.height = 10
ws_dash.add_chart(chart_line, "B10")

# 2. Top Products Bar Chart
chart_bar = BarChart()
chart_bar.type = "col"
chart_bar.style = 10
chart_bar.title = "Top 5 Products by Revenue"
chart_bar.y_axis.title = "Revenue"

prod_start_row = len(monthly_df) + 4
data_bar = Reference(ws_data, min_col=2, min_row=prod_start_row, max_row=prod_start_row+len(top_products_df))
cats_bar = Reference(ws_data, min_col=1, min_row=prod_start_row+1, max_row=prod_start_row+len(top_products_df))
chart_bar.add_data(data_bar, titles_from_data=True)
chart_bar.set_categories(cats_bar)
chart_bar.width = 16
chart_bar.height = 10
ws_dash.add_chart(chart_bar, "I10")

# Adjust column widths
for ws in [ws_dash, ws_data]:
    ws.views.sheetView[0].showGridLines = True
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

# Save output workbook
wb.save(output_excel_path)
print(f"Excel Dashboard successfully created at: {output_excel_path}")